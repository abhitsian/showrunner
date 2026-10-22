#!/usr/bin/env python3
"""sr: the Showrunner command line.

  sr new <name> [--from demo|talk|interview|training|sales|review]   start a show in ~/Showrunner from an example
  sr import <deck.pptx> [--pdf deck.pdf] [-o out.show.md]            a show from a deck: one item per slide, notes as the talk track
  sr from-autopilot <run.json> [-o out.show.json]                   convert a Handle Autopilot run of show
  sr check <show>                                                   parse it, list the items, run the preflight checks
  sr open <show>                                                    open it in Showrunner
  sr state | next | prev | go N | primary | toggle | beat | start | pause | reload    drive the open show
"""
import json
import os
import re
import shutil
import subprocess
import sys
import urllib.request

HERE = os.path.dirname(os.path.realpath(__file__))
APP = os.path.expanduser("~/Applications/Showrunner.app")
SHOWS = os.path.expanduser("~/Showrunner")
REMOTE = "http://127.0.0.1:4977/"


def die(msg):
    print(f"sr: {msg}", file=sys.stderr)
    sys.exit(1)


def arg(flag, default=None):
    if flag in sys.argv:
        k = sys.argv.index(flag)
        if k + 1 < len(sys.argv):
            return sys.argv[k + 1]
    return default


def cmd_new(name):
    kind = arg("--from", "demo")
    src = os.path.join(HERE, "examples", f"{kind}.show.md")
    if not os.path.exists(src):
        die(f"no example called {kind}; have: " + ", ".join(sorted(f.split('.')[0] for f in os.listdir(os.path.join(HERE, 'examples')) if f.endswith('.show.md'))))
    os.makedirs(SHOWS, exist_ok=True)
    dst = os.path.join(SHOWS, f"{name}.show.md")
    if os.path.exists(dst):
        die(f"{dst} exists")
    shutil.copy(src, dst)
    print(dst)
    subprocess.run(["open", "-a", APP, dst])


def clean(text):
    return re.sub(r"[ \t]+\n", "\n", text or "").strip()


def cmd_import(deck):
    try:
        from pptx import Presentation
    except ImportError:
        die("needs python-pptx: pip3 install python-pptx")
    deck = os.path.abspath(os.path.expanduser(deck))
    base = os.path.splitext(os.path.basename(deck))[0]
    out = os.path.abspath(arg("-o", os.path.join(os.path.dirname(deck), base + ".show.md")))
    imgdir = os.path.join(os.path.dirname(out), base + " slides")
    pdf = arg("--pdf")
    prs = Presentation(deck)
    slides = []
    for n, s in enumerate(prs.slides, 1):
        hidden = s._element.get("show") == "0"
        title = (s.shapes.title.text if s.shapes.title is not None else "").strip()
        if not title:
            texts = [sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            title = texts[0].split("\n")[0][:80] if texts else f"Slide {n}"
        notes = clean(s.notes_slide.notes_text_frame.text) if s.has_notes_slide else ""
        slides.append({"n": n, "title": " ".join(title.split()), "notes": notes, "hidden": hidden})

    # slide images: from a PDF export (hidden slides are left out of PDFs), else LibreOffice if installed
    images = {}
    if not pdf and shutil.which("soffice"):
        tmp = os.path.join(imgdir, "_pdf")
        os.makedirs(tmp, exist_ok=True)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, deck], capture_output=True)
        cand = os.path.join(tmp, base + ".pdf")
        pdf = cand if os.path.exists(cand) else None
    if pdf and shutil.which("pdftoppm"):
        os.makedirs(imgdir, exist_ok=True)
        shown = [s for s in slides if not s["hidden"]]
        for page, s in enumerate(shown, 1):
            stem = os.path.join(imgdir, f"slide-{s['n']:02d}")
            subprocess.run(["pdftoppm", "-f", str(page), "-l", str(page), "-png", "-scale-to-x", "1280", "-scale-to-y", "-1", "-singlefile", pdf, stem], check=False)
            if os.path.exists(stem + ".png"):
                images[s["n"]] = os.path.relpath(stem + ".png", os.path.dirname(out))

    lines = [f"title: {base}", "", f"# {base}", ""]
    for s in slides:
        if s["hidden"]:
            continue
        lines.append(f"## [slide {s['n']}] {s['title']}")
        if s["n"] in images:
            lines.append(f"@image {images[s['n']]}")
        lines.append(f"> Slide {s['n']}.")
        lines.append("")
        lines.append(s["notes"] or "(no speaker notes)")
        lines.append("")
    with open(out, "w") as f:
        f.write("\n".join(lines))
    hidden = sum(s["hidden"] for s in slides)
    print(f"{out}: {len(slides) - hidden} slides" + (f", {hidden} hidden left out" if hidden else "") +
          (f", {len(images)} images" if images else ", no images (pass --pdf deck.pdf, or install LibreOffice)"))


def cmd_from_autopilot(path):
    run = json.load(open(os.path.expanduser(path)))
    out = arg("-o", os.path.splitext(os.path.expanduser(path))[0] + ".show.json")
    kinds = {"tab": "demo", "slide": "slide", "talk": "talk"}
    sections = []
    for s in run.get("sections", []):
        items = []
        for it in s.get("items", []):
            n = {"kind": kinds.get(it.get("kind"), "talk"), "title": it.get("title", "")}
            for k in ("cue", "say", "note", "image", "slide"):
                if it.get(k):
                    n[k] = it[k]
            if it.get("flow") is not None:
                n["go"] = [{"handleFlow": it["flow"]}]
            items.append(n)
        sections.append({"title": s.get("title", ""), "items": items})
    show = {"title": run.get("title", ""), "sections": sections, "preflight": [{"handle": True}]}
    with open(out, "w") as f:
        json.dump(show, f, ensure_ascii=False, indent=2)
    print(out)


def remote(path):
    try:
        with urllib.request.urlopen(REMOTE + path, timeout=5) as r:
            print(r.read().decode())
    except Exception as e:
        die(f"Showrunner isn't answering on {REMOTE} ({e}). Is it open?")


def main():
    a = sys.argv[1:]
    if not a or a[0] in ("-h", "--help", "help"):
        print(__doc__)
        return
    c = a[0]
    if c == "new" and len(a) > 1:
        cmd_new(a[1])
    elif c == "import" and len(a) > 1:
        cmd_import(a[1])
    elif c == "from-autopilot" and len(a) > 1:
        cmd_from_autopilot(a[1])
    elif c == "check" and len(a) > 1:
        sys.exit(subprocess.run([os.path.join(APP, "Contents/MacOS/Showrunner"), "--check", os.path.abspath(os.path.expanduser(a[1]))]).returncode)
    elif c == "open" and len(a) > 1:
        subprocess.run(["open", "-a", APP, os.path.abspath(os.path.expanduser(a[1]))])
    elif c == "go" and len(a) > 1:
        remote(f"go/{a[1]}")
    elif c in ("state", "next", "prev", "back", "primary", "toggle", "start", "pause", "reload"):
        remote(c)
    elif c == "beat":
        remote("beat/prev" if len(a) > 1 and a[1] == "prev" else "beat")
    else:
        die(f"unknown command: {' '.join(a)}. Try sr --help")


if __name__ == "__main__":
    main()
